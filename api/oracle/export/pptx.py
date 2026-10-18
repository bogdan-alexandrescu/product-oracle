"""The deck as PPTX (O1f): the same slides as Present mode and the deck PDF.

A 16:9 presentation with a title slide (the version banner), one slide per brief
section and the Sources and data points slides. Each slide's body is built from the
section's markdown: paragraphs, bullet and numbered lists, minor headings and tables
(as native PowerPoint tables); presenter notes go into the slide's notes. Colours and
fonts come from the document's theme tokens; every slide carries the version banner
in its footer.
"""

import re
from dataclasses import dataclass, field
from io import BytesIO

from markdown_it import MarkdownIt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from oracle.docs.parser import CITATION_STRIP_RE
from oracle.export.deck import Deck, Slide

WIDTH, HEIGHT = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.8)
BODY_TOP = Inches(1.55)
FOOTER_TOP = Inches(6.95)
_md = MarkdownIt("commonmark", {"html": False}).enable(["table", "strikethrough"])
_HEX = re.compile(r"^#([0-9a-fA-F]{3}|[0-9a-fA-F]{6})$")


@dataclass
class Run:
    text: str
    bold: bool = False
    italic: bool = False
    mono: bool = False


@dataclass
class Para:
    runs: list[Run]
    level: int = 0
    bullet: str | None = None
    heading: bool = False


@dataclass
class Table:
    rows: list[list[str]]


@dataclass
class Blocks:
    items: list[Para | Table] = field(default_factory=list)


def _colour(tokens: dict, name: str, default: str) -> RGBColor:
    value = (tokens.get("colors") or {}).get(name) or default
    if not isinstance(value, str) or not _HEX.match(value):
        value = default
    h = value.lstrip("#")
    if len(h) == 3:
        h = "".join(c * 2 for c in h)
    return RGBColor.from_string(h.upper())


def _font(tokens: dict, name: str) -> str:
    stack = (tokens.get("fonts") or {}).get(name) or ""
    first = stack.split(",")[0].strip().strip("'\"")
    return first or "DejaVu Sans"


def _runs(children) -> list[Run]:
    runs: list[Run] = []
    bold = italic = False
    for t in children or []:
        if t.type == "text":
            runs.append(Run(t.content, bold, italic))
        elif t.type == "code_inline":
            runs.append(Run(t.content, bold, italic, mono=True))
        elif t.type in ("softbreak", "hardbreak"):
            runs.append(Run(" ", bold, italic))
        elif t.type == "strong_open":
            bold = True
        elif t.type == "strong_close":
            bold = False
        elif t.type == "em_open":
            italic = True
        elif t.type == "em_close":
            italic = False
    return runs


def parse_blocks(markdown: str) -> Blocks:
    """Slide body markdown as paragraphs and tables."""
    out = Blocks()
    tokens = _md.parse(CITATION_STRIP_RE.sub("", markdown or ""))
    lists: list[list] = []  # stack of [kind, counter]
    table: list[list[str]] | None = None
    row: list[str] | None = None
    heading = False
    for t in tokens:
        if t.type in ("bullet_list_open", "ordered_list_open"):
            lists.append(
                ["ul" if t.type == "bullet_list_open" else "ol", int(t.attrGet("start") or 1)]
            )
        elif t.type in ("bullet_list_close", "ordered_list_close"):
            lists.pop()
        elif t.type == "heading_open":
            heading = True
        elif t.type == "heading_close":
            heading = False
        elif t.type == "table_open":
            table = []
        elif t.type == "table_close":
            if table:
                out.items.append(Table(table))
            table = None
        elif t.type == "tr_open":
            row = []
        elif t.type == "tr_close" and table is not None and row is not None:
            table.append(row)
            row = None
        elif t.type == "inline":
            runs = _runs(t.children)
            if row is not None:
                row.append("".join(r.text for r in runs).strip())
                continue
            if not "".join(r.text for r in runs).strip():
                continue
            bullet = None
            if lists:
                kind = lists[-1]
                if kind[0] == "ul":
                    bullet = "•"
                else:
                    bullet = f"{kind[1]}."
                    kind[1] += 1
            out.items.append(
                Para(runs, level=max(len(lists) - 1, 0), bullet=bullet, heading=heading)
            )
    return out


def _text_height(paras: list[Para], size: int) -> int:
    # About 95 characters per line at 20pt across the body width.
    per_line = int(95 * 20 / size)
    lines = sum(max(1, -(-len("".join(r.text for r in p.runs)) // per_line)) for p in paras)
    return Emu(int(lines * Pt(size) * 1.45 + Pt(size) * 0.6 * len(paras)))


def _write_runs(paragraph, runs: list[Run], size: int, font: str, mono: str, colour) -> None:
    for r in runs:
        run = paragraph.add_run()
        run.text = r.text
        run.font.size = Pt(size)
        run.font.bold = r.bold or None
        run.font.italic = r.italic or None
        run.font.name = mono if r.mono else font
        run.font.color.rgb = colour


def _textbox(slide, left, top, width, height, *, wrap=True):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = wrap
    frame.margin_left = frame.margin_right = 0
    return frame


def _add_body(slide, s: Slide, deck: Deck, style: dict) -> None:
    blocks = parse_blocks(s.body_md)
    size = 16 if s.dense else 20
    top = BODY_TOP
    width = WIDTH - 2 * MARGIN
    pending: list[Para] = []

    def flush_text() -> None:
        nonlocal top, pending
        if not pending:
            return
        height = _text_height(pending, size)
        frame = _textbox(slide, MARGIN, top, width, height)
        for i, p in enumerate(pending):
            para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            para.level = min(p.level, 4)
            para.space_after = Pt(size * 0.4)
            runs = list(p.runs)
            if p.bullet:
                runs.insert(0, Run(f"{p.bullet} "))
            if p.heading:
                runs = [Run(r.text, True, r.italic, r.mono) for r in runs]
            _write_runs(para, runs, size, style["body"], style["mono"], style["text"])
        top += height
        pending = []

    for item in blocks.items:
        if isinstance(item, Para):
            pending.append(item)
            continue
        flush_text()
        cols = max(len(r) for r in item.rows)
        row_h = Inches(0.42 if not s.dense else 0.34)
        height = row_h * len(item.rows)
        shape = slide.shapes.add_table(len(item.rows), cols, MARGIN, top, width, height)
        tbl = shape.table
        for ri, cells in enumerate(item.rows):
            for ci in range(cols):
                cell = tbl.cell(ri, ci)
                cell.text = ""
                para = cell.text_frame.paragraphs[0]
                text = cells[ci] if ci < len(cells) else ""
                _write_runs(
                    para,
                    [Run(text, bold=ri == 0)],
                    max(size - 6, 11),
                    style["body"],
                    style["mono"],
                    style["heading"] if ri == 0 else style["text"],
                )
        top += height + Inches(0.15)
    flush_text()


def deck_to_pptx(deck: Deck) -> bytes:
    tokens = deck.tokens or {}
    heading_font = _font(tokens, "heading")
    style = {
        "body": _font(tokens, "body"),
        "mono": _font(tokens, "mono"),
        "text": _colour(tokens, "text", "#1f2933"),
        "heading": _colour(tokens, "heading", "#111827"),
        "muted": _colour(tokens, "muted", "#5b6472"),
        "accent": _colour(tokens, "accent", "#2f5bd3"),
        "background": _colour(tokens, "background", "#ffffff"),
    }
    prs = Presentation()
    prs.slide_width, prs.slide_height = WIDTH, HEIGHT
    blank = prs.slide_layouts[6]
    prs.core_properties.title = f"{deck.title} — v{deck.number} (Deck)"
    prs.core_properties.subject = deck.banner
    total = len(deck.slides)

    for s in deck.slides:
        slide = prs.slides.add_slide(blank)
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = style["background"]
        width = WIDTH - 2 * MARGIN
        if s.kind == "title":
            kicker = _textbox(slide, MARGIN, Inches(2.2), width, Inches(0.5))
            _write_runs(
                kicker.paragraphs[0],
                [Run("BRIEF", bold=True)],
                16,
                heading_font,
                style["mono"],
                style["accent"],
            )
            title = _textbox(slide, MARGIN, Inches(2.75), width, Inches(1.6))
            _write_runs(
                title.paragraphs[0],
                [Run(deck.title, bold=True)],
                44,
                heading_font,
                style["mono"],
                style["heading"],
            )
            sub = _textbox(slide, MARGIN, Inches(4.4), width, Inches(0.8))
            status = deck.status.replace("_", " ").capitalize()
            line = f"Version {deck.number} · {status} · as of {deck.as_of}"
            _write_runs(
                sub.paragraphs[0], [Run(line)], 20, style["body"], style["mono"], style["muted"]
            )
        else:
            title = _textbox(slide, MARGIN, Inches(0.55), width, Inches(0.9))
            title.vertical_anchor = MSO_ANCHOR.TOP
            size = 28 if s.dense else 32
            _write_runs(
                title.paragraphs[0],
                [Run(s.heading, bold=True)],
                size,
                heading_font,
                style["mono"],
                style["heading"],
            )
            if s.kind == "sources":
                frame = _textbox(slide, MARGIN, BODY_TOP, width, Inches(5))
                for i, (number, src_title, kind, retrieved) in enumerate(s.sources):
                    para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
                    para.space_after = Pt(6)
                    extra = f"  {kind}" + (f", retrieved {retrieved}" if retrieved else "")
                    _write_runs(
                        para,
                        [Run(f"{number}. {src_title}")],
                        16,
                        style["body"],
                        style["mono"],
                        style["text"],
                    )
                    _write_runs(
                        para, [Run(extra)], 12, style["body"], style["mono"], style["muted"]
                    )
            else:
                _add_body(slide, s, deck, style)
        footer = _textbox(slide, MARGIN, FOOTER_TOP, width, Inches(0.35), wrap=False)
        para = footer.paragraphs[0]
        _write_runs(
            para,
            [Run(f"{deck.banner}    {s.index} / {total}")],
            10,
            style["body"],
            style["mono"],
            style["muted"],
        )
        para.alignment = PP_ALIGN.LEFT
        if s.notes:
            slide.notes_slide.notes_text_frame.text = s.notes

    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()
