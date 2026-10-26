"""O5 lane 4: company branding (spec-O5 section 4, acceptance criteria 14 and 15, owner
answer O5-Q8).

Card: protects=the ``branding`` setting (display name, logo, dark logo, accent, public
disclaimer, the "Made with Product Oracle" switch, a revision that bumps on every change)
read by every member and changed by owners only; a logo is a PNG, JPEG or WebP of at most
512 KB and 2,000 x 2,000 px, re-encoded to PNG, stored in ``brand_assets`` and served at
``/uploads/brand/<sha256>.png`` with a safe content type; every HTML/PDF export, the deck and
the PPTX title slide carry the logo as the one ``data:`` image of the header, the accent
replaces the theme's accent, and no export carries Oracle's mark;
fails_when=an SVG, a script disguised as a PNG, an oversized or a too-large image is
stored, a logo is served as anything but image/png or to another company, a non-owner or a
run changes branding, a too-light accent is saved, the export header lacks the logo, a
body ``data:`` image renders, or an export shows "Made with Product Oracle";
no_claim=the shared page (lanes 3 and 5) and the PDF's pixels (WeasyPrint is not installed
in every lane; the PDF test is skipped without it).
"""

import base64
import hashlib
import io
import re
import uuid

import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader

from oracle.auth.deps import RUN_TOKEN_HEADER
from oracle.runs.models import Run
from oracle.runtokens import mint_run_token
from oracle.tenancy import tenant_session
from tests.baseline_helpers import Ctx, make_ctx
from tests.conftest import seed_company
from tests.coverage_helpers import run_headers
from tests.test_brief import DRAFT, _import, requires_pdf

TOO_LIGHT = "This colour is too light to read on white. Pick a darker one."
NOT_AN_IMAGE = "Upload a PNG, JPEG or WebP"
DEFAULT_DISCLAIMER = (
    "This roadmap shows our current plans. It is not a commitment, and plans and timing can change."
)
MARK = "Made with Product Oracle"


def _image(fmt: str = "PNG", size: tuple[int, int] = (320, 80), colour=(3, 105, 161)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", size, colour).save(buf, format=fmt)
    return buf.getvalue()


def _noisy_png(min_bytes: int) -> bytes:
    """A PNG over ``min_bytes`` (random pixels do not compress)."""
    import os

    side = 1100
    img = Image.frombytes("RGB", (side, side), os.urandom(side * side * 3))
    buf = io.BytesIO()
    img.save(buf, format="PNG", compress_level=0)
    data = buf.getvalue()
    assert len(data) > min_bytes
    return data


SVG = (
    b'<svg xmlns="http://www.w3.org/2000/svg" width="10" height="10">'
    b"<script>alert(1)</script></svg>"
)


@pytest.fixture
async def ctx() -> Ctx:
    ctx = await make_ctx("Example Labs")
    await seed_company(ctx.tenant_id)
    return ctx


async def _scoped_run_token(ctx: Ctx) -> dict[str, str]:
    """A run's own token (scoped, as a run's agents call the API)."""
    async with tenant_session(ctx.tenant_id) as s:
        run = Run(flow="baseline_to_prd", status="running", created_by=ctx.users["owner"].id)
        s.add(run)
        await s.flush()
        raw = await mint_run_token(s, run.id, ["read:baseline"], ttl_seconds=600)
        await s.commit()
    return {**ctx.as_("owner"), RUN_TOKEN_HEADER: raw}


async def _upload(
    client, headers, data: bytes, *, name="logo.png", ctype="image/png", variant=None
):
    params = {"variant": variant} if variant else None
    return await client.post(
        "/v1/branding/logo", params=params, files={"file": (name, data, ctype)}, headers=headers
    )


def _body(display_name="Example Labs", accent=None, disclaimer=DEFAULT_DISCLAIMER, mark=True):
    return {
        "display_name": display_name,
        "accent": accent,
        "public_disclaimer": disclaimer,
        "oracle_mark": mark,
    }


# --- the setting -----------------------------------------------------------------------


async def test_every_member_reads_the_defaults_and_only_owners_may_edit(client, ctx):
    for who in ("viewer", "commenter", "approver", "owner", "token"):
        resp = await client.get("/v1/branding", headers=ctx.as_(who))
        assert resp.status_code == 200, (who, resp.text)
        out = resp.json()
        assert out["display_name"] == "Example Labs"  # the company's name
        assert out["logo"] is None and out["logo_dark"] is None
        assert out["accent"] is None
        assert out["public_disclaimer"] == DEFAULT_DISCLAIMER
        assert out["oracle_mark"] is True  # O5-Q8: on until an owner switches it off
        assert out["revision"] == 0
        assert out["can_edit"] is (who in ("owner", "token")), who


async def test_a_run_reads_branding_but_never_changes_it(client, ctx):
    bound = await run_headers(ctx, "owner")
    assert (await client.get("/v1/branding", headers=bound)).status_code == 200
    scoped = await _scoped_run_token(ctx)
    assert (await client.get("/v1/branding", headers=scoped)).status_code == 403
    for headers in (bound, scoped):
        resp = await client.put("/v1/branding", json=_body("Run Co"), headers=headers)
        assert resp.status_code == 403, resp.text
        assert (await _upload(client, headers, _image())).status_code == 403
        resp = await client.delete("/v1/branding/logo", headers=headers)
        assert resp.status_code == 403
    assert (await client.get("/v1/branding", headers=ctx.as_("owner"))).json()["revision"] == 0


async def test_only_owners_change_branding(client, ctx):
    for who in ("viewer", "commenter", "approver"):
        resp = await client.put("/v1/branding", json=_body("Not Mine"), headers=ctx.as_(who))
        assert resp.status_code == 403, who
        assert (await _upload(client, ctx.as_(who), _image())).status_code == 403, who
        assert (await client.delete("/v1/branding/logo", headers=ctx.as_(who))).status_code == 403
    resp = await client.put(
        "/v1/branding", json=_body("Example Labs Ltd"), headers=ctx.as_("owner")
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["display_name"] == "Example Labs Ltd"
    assert resp.json()["revision"] == 1
    # An owner's API token is an owner too.
    resp = await client.put(
        "/v1/branding", json=_body("Example Labs", mark=False), headers=ctx.as_("token")
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["oracle_mark"] is False and resp.json()["revision"] == 2
    viewer = (await client.get("/v1/branding", headers=ctx.as_("viewer"))).json()
    assert viewer["oracle_mark"] is False and viewer["display_name"] == "Example Labs"


async def test_accent_is_contrast_checked_and_values_are_validated(client, ctx):
    owner = ctx.as_("owner")
    resp = await client.put("/v1/branding", json=_body(accent="#e0f2fe"), headers=owner)
    assert resp.status_code == 422
    assert resp.json()["detail"] == TOO_LIGHT
    resp = await client.put("/v1/branding", json=_body(accent="#0369a1"), headers=owner)
    assert resp.status_code == 200, resp.text
    assert resp.json()["accent"] == "#0369a1"
    for bad in (
        _body(accent="blue"),
        _body(accent="#0369a1; background:url(x)"),
        _body(display_name="  "),
        _body(display_name="x" * 81),
        _body(disclaimer=""),
        _body(disclaimer="x" * 501),
    ):
        assert (await client.put("/v1/branding", json=bad, headers=owner)).status_code == 422, bad
    assert (await client.get("/v1/branding", headers=owner)).json()["accent"] == "#0369a1"


async def test_the_generic_settings_route_does_not_write_branding(client, ctx):
    owner = ctx.as_("owner")
    keys = [s["key"] for s in (await client.get("/v1/settings", headers=owner)).json()]
    assert "branding" in keys
    resp = await client.put(
        "/v1/settings/branding", json={"value": {"display_name": "X"}}, headers=owner
    )
    assert resp.status_code == 422
    assert "Branding" in resp.json()["detail"]


# --- the logo --------------------------------------------------------------------------


async def test_only_png_jpeg_or_webp_logos_within_the_limits_are_stored(client, ctx):
    owner = ctx.as_("owner")
    cases = [
        (SVG, "logo.svg", "image/svg+xml"),
        (SVG, "logo.png", "image/png"),  # a script dressed as a PNG
        (b"<html><script>alert(1)</script></html>", "logo.png", "image/png"),
        (_image("GIF"), "logo.gif", "image/gif"),
        (_noisy_png(3 * 1024 * 1024 - 1), "big.png", "image/png"),  # over 512 KB
        (_image(size=(2001, 10)), "wide.png", "image/png"),  # over 2,000 px
        (b"", "empty.png", "image/png"),
    ]
    for data, name, ctype in cases:
        resp = await _upload(client, owner, data, name=name, ctype=ctype)
        assert resp.status_code == 422, (name, len(data), resp.text)
    resp = await _upload(client, owner, SVG, name="logo.svg", ctype="image/svg+xml")
    assert NOT_AN_IMAGE in resp.json()["detail"]
    out = (await client.get("/v1/branding", headers=owner)).json()
    assert out["logo"] is None and out["revision"] == 0
    async with tenant_session(ctx.tenant_id) as s:
        from sqlalchemy import func, select

        from oracle.branding import BrandAsset

        assert await s.scalar(select(func.count()).select_from(BrandAsset)) == 0


async def test_a_logo_is_re_encoded_to_png_and_served_safely_to_members_only(client, ctx):
    owner = ctx.as_("owner")
    resp = await _upload(client, owner, _image("JPEG"), name="logo.jpg", ctype="image/jpeg")
    assert resp.status_code == 200, resp.text
    logo = resp.json()["logo"]
    assert logo["type"] == "image/png"
    assert (logo["width"], logo["height"]) == (320, 80)
    assert logo["url"] == f"/uploads/brand/{logo['sha256']}.png"
    assert resp.json()["revision"] == 1
    for who in ("viewer", "commenter", "approver", "owner", "token"):
        served = await client.get(logo["url"], headers=ctx.as_(who))
        assert served.status_code == 200, who
        assert served.headers["content-type"] == "image/png"
        assert served.headers["x-content-type-options"] == "nosniff"
        assert "default-src 'none'" in served.headers["content-security-policy"]
        assert hashlib.sha256(served.content).hexdigest() == logo["sha256"]
        assert served.content.startswith(b"\x89PNG\r\n\x1a\n")  # re-encoded, not the JPEG
    # A scoped run token does not read company files.
    assert (await client.get(logo["url"], headers=await _scoped_run_token(ctx))).status_code == 403
    # Nobody signed in, an unknown or malformed name: nothing.
    assert (await client.get(logo["url"])).status_code in (401, 403)
    assert (await client.get(f"/uploads/brand/{'0' * 64}.png", headers=owner)).status_code == 404
    assert (await client.get("/uploads/brand/..%2f..%2fetc.png", headers=owner)).status_code == 404
    # Another company never reads it.
    other = await make_ctx("Northwind")
    assert (await client.get(logo["url"], headers=other.as_("owner"))).status_code == 404


async def test_a_dark_variant_and_removing_logos(client, ctx):
    owner = ctx.as_("owner")
    light = (await _upload(client, owner, _image())).json()
    dark = await _upload(
        client,
        owner,
        _image("WEBP", colour=(250, 250, 250)),
        name="d.webp",
        ctype="image/webp",
        variant="dark",
    )
    assert dark.status_code == 200, dark.text
    out = dark.json()
    assert out["logo"] == light["logo"] and out["logo_dark"]["type"] == "image/png"
    assert out["revision"] == 2
    assert (await _upload(client, owner, _image(), variant="purple")).status_code == 422
    resp = await client.delete("/v1/branding/logo", params={"variant": "dark"}, headers=owner)
    assert resp.status_code == 200
    assert resp.json()["logo_dark"] is None and resp.json()["logo"] == light["logo"]
    resp = await client.delete("/v1/branding/logo", headers=owner)
    assert resp.json()["logo"] is None and resp.json()["revision"] == 4
    assert (await client.get(light["logo"]["url"], headers=owner)).status_code == 404


# --- exports ---------------------------------------------------------------------------


def _data_png(html: str) -> list[str]:
    return re.findall(r'<img[^>]*src="(data:image/png;base64,[A-Za-z0-9+/=]+)"', html)


@pytest.fixture
async def branded(client, ctx):
    owner = ctx.as_("owner")
    assert (await _import(client, ctx.as_("token"), DRAFT)).status_code == 201
    png = _image()
    logo = (await _upload(client, owner, png)).json()["logo"]
    resp = await client.put("/v1/branding", json=_body(accent="#0369a1"), headers=owner)
    assert resp.status_code == 200
    return ctx, logo


async def test_every_html_export_header_holds_the_logo_as_its_one_data_image(client, branded):
    ctx, logo = branded
    base = "/v1/documents/checkout/versions/1/export"
    for query in ("format=html", "layer=brief&format=html", "layer=deck&format=html"):
        resp = await client.get(f"{base}?{query}", headers=ctx.as_("viewer"))
        assert resp.status_code == 200, (query, resp.text)
        html = resp.text
        images = _data_png(html)
        assert len(images) == 1, query
        raw = base64.b64decode(images[0].split(",", 1)[1])
        assert hashlib.sha256(raw).hexdigest() == logo["sha256"]
        assert 'class="brand-logo"' in html
        assert "Example Labs" in html
        assert "#0369a1" in html  # the accent replaces the theme's
        assert MARK not in html  # O5-Q8: never on exports


async def test_an_unthemed_export_carries_the_logo_and_accent_too(client, ctx):
    md = "---\ntitle: Plain notes\n---\n\n# Plain notes {#title}\n\n## Background {#bg}\nText.\n"
    assert (await _import(client, ctx.as_("token"), md, doc="plain-notes")).status_code == 201
    owner = ctx.as_("owner")
    await _upload(client, owner, _image())
    await client.put("/v1/branding", json=_body(accent="#0369a1"), headers=owner)
    from oracle.export.render import load_export_version, render_html

    async with tenant_session(ctx.tenant_id) as s:
        ev = await load_export_version(s, "plain-notes", 1)
        ev.theme = None  # the plain PRD template, as before themes
        html = render_html(ev)
    assert len(_data_png(html)) == 1 and "#0369a1" in html and "Example Labs" in html


async def test_a_data_image_in_a_body_still_does_not_render(client, branded):
    ctx, _ = branded
    md = DRAFT.replace(
        "We will cut checkout to one page for returning buyers.",
        "We will cut checkout. ![pixel](data:image/png;base64,iVBORw0KGgo=)",
    )
    assert (await _import(client, ctx.as_("token"), md)).status_code == 201
    resp = await client.get(
        "/v1/documents/checkout/versions/2/export?format=html", headers=ctx.as_("viewer")
    )
    assert len(_data_png(resp.text)) == 1  # the header's logo only
    assert "iVBORw0KGgo=" not in resp.text.split("</header>", 1)[1]


async def test_without_a_logo_the_header_shows_the_company_name(client, ctx):
    assert (await _import(client, ctx.as_("token"), DRAFT)).status_code == 201
    resp = await client.get(
        "/v1/documents/checkout/versions/1/export?format=html", headers=ctx.as_("viewer")
    )
    assert _data_png(resp.text) == []
    assert 'class="brand-name"' in resp.text and "Example Labs" in resp.text


async def test_the_pptx_and_present_title_slides_carry_the_logo(client, branded):
    ctx, logo = branded
    viewer = ctx.as_("viewer")
    resp = await client.get(
        "/v1/documents/checkout/versions/1/export?layer=deck&format=pptx", headers=viewer
    )
    assert resp.status_code == 200
    prs = Presentation(io.BytesIO(resp.content))
    title = prs.slides[0]
    pictures = [sh for sh in title.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pictures) == 1
    assert hashlib.sha256(pictures[0].image.blob).hexdigest() == logo["sha256"]
    assert all(
        sh.shape_type != MSO_SHAPE_TYPE.PICTURE
        for slide in list(prs.slides)[1:]
        for sh in slide.shapes
    )
    texts = " ".join(sh.text_frame.text for sh in title.shapes if sh.has_text_frame)
    assert "Example Labs" in texts and MARK not in texts

    deck = (await client.get("/v1/documents/checkout/versions/1/deck", headers=viewer)).json()
    assert len(_data_png(deck["slides"][0]["html"])) == 1
    assert "#0369a1" in deck["css"]
    # A new logo shows at once in Present mode (the kept deck is keyed by the brand).
    await _upload(client, ctx.as_("owner"), _image(colour=(10, 10, 10)))
    again = (await client.get("/v1/documents/checkout/versions/1/deck", headers=viewer)).json()
    assert _data_png(again["slides"][0]["html"]) != _data_png(deck["slides"][0]["html"])


@requires_pdf
async def test_the_pdf_shows_the_logo(client, branded):
    ctx, _ = branded
    resp = await client.get(
        "/v1/documents/checkout/versions/1/export?format=pdf", headers=ctx.as_("viewer")
    )
    assert resp.status_code == 200
    page = PdfReader(io.BytesIO(resp.content)).pages[0]
    assert len(page.images) >= 1


# --- the shared page's header (lanes 3 and 5 render it) --------------------------------


async def test_the_shared_header_has_the_mark_unless_an_owner_switches_it_off(client, branded):
    from oracle.branding import load_brand, render_brand_header

    ctx, _ = branded
    async with tenant_session(ctx.tenant_id) as s:
        brand = await load_brand(s)
    shared = render_brand_header(brand, shared=True)
    assert MARK in shared and len(_data_png(shared)) == 1
    assert MARK not in render_brand_header(brand, shared=False)
    await client.put(
        "/v1/branding", json=_body(accent="#0369a1", mark=False), headers=ctx.as_("owner")
    )
    async with tenant_session(ctx.tenant_id) as s:
        brand = await load_brand(s)
    assert MARK not in render_brand_header(brand, shared=True)
    assert brand.revision == 3 and isinstance(brand.tenant_id, uuid.UUID)


# --- #627 review: caches, resets and failures --------------------------------------------


async def test_the_render_key_changes_with_the_mark_and_the_disclaimer(client, branded):
    from oracle.branding import load_brand

    ctx, _ = branded
    owner = ctx.as_("owner")
    async with tenant_session(ctx.tenant_id) as s:
        before = (await load_brand(s)).key
    await client.put("/v1/branding", json=_body(accent="#0369a1", mark=False), headers=owner)
    async with tenant_session(ctx.tenant_id) as s:
        no_mark = (await load_brand(s)).key
    assert no_mark != before
    body = _body(accent="#0369a1", mark=False, disclaimer="Plans change.")
    await client.put("/v1/branding", json=body, headers=owner)
    async with tenant_session(ctx.tenant_id) as s:
        assert (await load_brand(s)).key != no_mark


async def test_a_reset_keeps_the_revision_counting_and_drops_the_logos(client, branded):
    from sqlalchemy import func, select

    from oracle.branding import BrandAsset

    ctx, logo = branded
    owner = ctx.as_("owner")
    assert (await client.get("/v1/branding", headers=owner)).json()["revision"] == 2
    resp = await client.put("/v1/settings/branding", json={"value": None}, headers=owner)
    assert resp.status_code == 200, resp.text
    out = (await client.get("/v1/branding", headers=owner)).json()
    assert out["revision"] == 3
    assert out["logo"] is None and out["accent"] is None and out["display_name"] == "Example Labs"
    assert (await client.get(logo["url"], headers=owner)).status_code == 404
    async with tenant_session(ctx.tenant_id) as s:
        assert await s.scalar(select(func.count()).select_from(BrandAsset)) == 0
    # The next change counts on from there.
    assert (await _upload(client, owner, _image())).json()["revision"] == 4


async def test_a_raw_branding_value_on_the_settings_route_is_refused_in_the_app_s_words(
    client, ctx
):
    """#654 review: the refusal names the screen a person uses, not the route or method."""
    from oracle import wording

    owner = ctx.as_("owner")
    resp = await client.put(
        "/v1/settings/branding", json={"value": {"accent": "#0369a1"}}, headers=owner
    )
    assert resp.status_code == 422, resp.text
    detail = resp.json()["detail"]
    assert "Settings › Company › Branding" in detail, detail
    assert "a null value here resets every default" in detail, detail
    assert not wording.api_words(detail), detail


async def test_the_same_logo_uploaded_twice_at_once_is_stored_once(client, ctx):
    """A second upload of one image, landing while the first is uncommitted, reuses the
    row rather than failing on the unique index."""
    import asyncio

    from sqlalchemy import func, select, text

    from oracle.branding import BrandAsset, reencode_logo

    owner = ctx.as_("owner")
    data = _image()
    png, width, height = reencode_logo(data)
    sha = hashlib.sha256(png).hexdigest()
    async with tenant_session(ctx.tenant_id) as first:
        first.add(
            BrandAsset(sha256=sha, content_type="image/png", bytes=png, width=width, height=height)
        )
        await first.flush()
        upload = asyncio.create_task(_upload(client, owner, data))
        for _ in range(100):  # until the upload waits on the uncommitted row
            await asyncio.sleep(0.05)
            async with tenant_session(ctx.tenant_id) as probe:
                waiting = await probe.scalar(
                    text(
                        "SELECT count(*) FROM pg_stat_activity WHERE wait_event_type = 'Lock' "
                        "AND query ILIKE 'INSERT INTO brand_assets%'"
                    )
                )
            if waiting or upload.done():
                break
        await first.commit()
    resp = await upload
    assert resp.status_code == 200, resp.text
    assert resp.json()["logo"]["sha256"] == sha
    async with tenant_session(ctx.tenant_id) as s:
        assert await s.scalar(select(func.count()).select_from(BrandAsset)) == 1
    assert (await client.get(resp.json()["logo"]["url"], headers=owner)).status_code == 200


async def test_an_image_mode_pillow_cannot_convert_is_refused_not_a_crash(client, ctx, monkeypatch):
    from PIL import Image as PILImage

    def refuse(self, *args, **kwargs):
        raise ValueError("conversion not supported")

    monkeypatch.setattr(PILImage.Image, "convert", refuse)
    buf = io.BytesIO()
    Image.new("1", (40, 10)).save(buf, format="PNG")  # mode "1" goes through convert
    resp = await _upload(client, ctx.as_("owner"), buf.getvalue())
    assert resp.status_code == 422, resp.text
    assert NOT_AN_IMAGE in resp.json()["detail"]
