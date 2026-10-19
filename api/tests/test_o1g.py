"""O1g: templates must mark a brief, company themes carry slide CSS, and presenter notes
come only from a ``notes`` block."""

import io
import time

import pytest
from pptx import Presentation
from pypdf import PdfReader
from sqlalchemy import select

from oracle.docs.models import Document
from oracle.export.render import load_export_version, render_html
from oracle.main import app
from oracle.mcp.tools import app_operations
from oracle.models import AuditEvent
from oracle.templates.builtins import BUILTIN_TEMPLATES, BUILTIN_THEMES, builtin_body
from oracle.templates.models import Template, TemplateVersion, Theme
from oracle.templates.validation import check_slide_css, check_template, check_theme_css
from oracle.tenancy import tenant_session
from tests.conftest import seed_company
from tests.test_brief import BRIEF_KEYS, DRAFT, _import, requires_pdf
from tests.test_priorities import _token
from tests.test_templates import THEME

BRIEFLESS = "# T {#title}\n\n## Problem {#problem}\n\n## Sources and data points {#sources}\n"
MARKED = "---\nbrief: [problem]\n---\n" + BRIEFLESS
FULL_HEADING = (
    "# T {#title}\n\n## Problem {#problem}\n\n# Full version {#full}\n\n"
    "## Detail {#detail}\n\n## Sources and data points {#sources}\n"
)
SLIDE_CSS = '.deck .slide-title::after { content: " SLIDECSS"; }'


# --- templates must mark a brief --------------------------------------------------------


def test_every_builtin_passes_with_its_brief():
    for t in BUILTIN_TEMPLATES:
        check = check_template(builtin_body(t))
        assert check.ok, (t.key, check.problems)
        assert check.brief_sections, t.key


async def test_validate_refuses_a_template_without_brief_and_says_how_to_fix_it(client, tenant_id):
    h = await _token(tenant_id, "viewer")
    out = (await client.post("/v1/templates/validate", json={"body": BRIEFLESS}, headers=h)).json()
    assert out["valid"] is False and out["warnings"] == [] and out["brief_sections"] == []
    [problem] = out["problems"]
    assert problem["code"] == "no_brief_sections"
    assert "brief: [section ids]" in problem["message"]
    assert "# Full version {#full}" in problem["message"]
    for body, brief in ((MARKED, ["problem"]), (FULL_HEADING, ["problem"])):
        out = (await client.post("/v1/templates/validate", json={"body": body}, headers=h)).json()
        assert out["valid"] and out["problems"] == [] and out["brief_sections"] == brief


async def test_create_and_edit_refuse_a_template_without_brief(client, tenant_id):
    h = await _token(tenant_id, "approver")
    await seed_company(tenant_id)
    payload = {"key": "plain", "name": "Plain", "kind": "prd", "body": BRIEFLESS}
    resp = await client.post("/v1/templates", json=payload, headers=h)
    assert resp.status_code == 422
    assert [p["code"] for p in resp.json()["problems"]] == ["no_brief_sections"]

    assert (
        await client.post("/v1/templates", json={**payload, "body": MARKED}, headers=h)
    ).status_code == 201
    resp = await client.patch("/v1/templates/plain", json={"body": BRIEFLESS}, headers=h)
    assert resp.status_code == 422
    assert [p["code"] for p in resp.json()["problems"]] == ["no_brief_sections"]
    t = (await client.get("/v1/templates/plain", headers=h)).json()
    assert t["latest_version"] == 1 and t["versions"][0]["warnings"] == []

    # A built-in edited to drop its brief is refused too.
    nnl = (await client.get("/v1/templates/now-next-later", headers=h)).json()
    body = nnl["versions"][-1]["body"]
    stripped = body.replace("brief: [", "old_brief: [").replace(
        "# Full version {#full}", "# Everything else {#rest}"
    )
    resp = await client.patch("/v1/templates/now-next-later", json={"body": stripped}, headers=h)
    assert resp.status_code == 422
    assert "no_brief_sections" in {p["code"] for p in resp.json()["problems"]}


async def test_a_stored_template_without_brief_keeps_working_until_its_next_edit(client, tenant_id):
    h = await _token(tenant_id, "approver")
    await seed_company(tenant_id)
    # A company template saved before O1g, with no brief markers.
    async with tenant_session(tenant_id) as s:
        t = Template(key="legacy", name="Legacy", kind="prd", builtin=False)
        s.add(t)
        await s.flush()
        s.add(
            TemplateVersion(
                template_id=t.id,
                version=1,
                body=BRIEFLESS,
                required_sections=["title", "problem", "sources"],
                created_by=t.id,
            )
        )
        await s.commit()
    t = (await client.get("/v1/templates/legacy", headers=h)).json()
    [v1] = t["versions"]
    assert v1["body"] == BRIEFLESS and v1["brief_sections"] == []
    assert [w["code"] for w in v1["warnings"]] == ["no_brief_sections"]

    # Documents drafted from it keep working: they open on the full document.
    started = await client.post(
        "/v1/drafts", json={"kind": "prd", "title": "Old style", "template": "legacy"}, headers=h
    )
    assert started.status_code == 201, started.text
    version = (
        await client.get(f"/v1/documents/{started.json()['slug']}/versions/1", headers=h)
    ).json()
    assert version["brief_source"] == "none"

    # Renaming is not a new version; a new body must add brief markers.
    renamed = await client.patch("/v1/templates/legacy", json={"name": "Legacy PRD"}, headers=h)
    assert renamed.status_code == 200 and renamed.json()["latest_version"] == 1
    edited = BRIEFLESS.replace("## Problem", "## The problem")
    resp = await client.patch("/v1/templates/legacy", json={"body": edited}, headers=h)
    assert resp.status_code == 422
    assert [p["code"] for p in resp.json()["problems"]] == ["no_brief_sections"]
    resp = await client.patch("/v1/templates/legacy", json={"body": MARKED}, headers=h)
    assert resp.status_code == 200 and resp.json()["latest_version"] == 2
    assert resp.json()["versions"][-1]["warnings"] == []
    # Version 1 is not rewritten.
    v1_again = (await client.get("/v1/templates/legacy/versions/1", headers=h)).json()
    assert v1_again["body"] == BRIEFLESS


async def test_template_roles_are_unchanged(client, tenant_id):
    viewer = await _token(tenant_id, "viewer")
    commenter = await _token(tenant_id, "commenter")
    payload = {"key": "x", "name": "X", "kind": "prd", "body": MARKED}
    for h in (viewer, commenter):
        assert (await client.post("/v1/templates", json=payload, headers=h)).status_code == 403
    # Validation is open to anyone signed in, and only to them.
    assert (
        await client.post("/v1/templates/validate", json={"body": BRIEFLESS}, headers=viewer)
    ).status_code == 200
    anonymous = await client.post("/v1/templates/validate", json={"body": BRIEFLESS})
    assert anonymous.status_code == 401


# --- slide CSS on company themes ----------------------------------------------------------

_IMAGE_SET = '.slide{background-image:image-set("https://evil.example/p.png" 1x)}'


@pytest.mark.parametrize(
    "bad",
    [
        "@import 'https://evil.example/x.css';",
        ".slide { background: url(https://evil.example/x.png); }",
        ".slide { color: red; } </style><script>alert(1)</script>",
        ".slide { color: red;",
        ".slide { content: '\\41'; }",
        '.slide { background-image: image-set("https://evil.example/p.png" 1x); }',
        '.slide { background-image: -webkit-image-set("https://evil.example/p.png" 1x); }',
        '.slide { background: cross-fade("https://evil.example/a.png", red 50%); }',
        '.slide { background: image("https://evil.example/p.png"); }',
        # Hidden from a check that strips comments or strings first; a browser reads them.
        '.a::before{content:"/*"} ' + _IMAGE_SET + ' .b::before{content:"*/"}',
        '.a::before{content:"x\n}\n' + _IMAGE_SET,
        ".a{background:url(data:x/*)} " + _IMAGE_SET + " .b{background:url(data:y*/)}",
    ],
)
async def test_slide_css_is_checked_like_the_theme_css(client, tenant_id, bad):
    h = await _token(tenant_id, "approver")
    resp = await client.post("/v1/themes", json={**THEME, "slide_css": bad}, headers=h)
    assert resp.status_code == 422
    problems = resp.json()["problems"]
    assert problems and all(p["code"] in ("unsafe_css", "invalid_css") for p in problems)
    assert all(p["message"].startswith("slide CSS: ") for p in problems)
    check = await client.post(
        "/v1/themes/validate", json={"tokens": THEME["tokens"], "slide_css": bad}, headers=h
    )
    assert check.json()["valid"] is False
    # The same text as the theme's CSS fails for the same reason.
    same = await client.post("/v1/themes/validate", json={"tokens": {}, "css": bad}, headers=h)
    assert {p["code"] for p in same.json()["problems"]} == {p["code"] for p in problems}

    assert (await client.post("/v1/themes", json=THEME, headers=h)).status_code == 201
    resp = await client.patch("/v1/themes/brand", json={"slide_css": bad}, headers=h)
    assert resp.status_code == 422
    assert (await client.get("/v1/themes/brand", headers=h)).json()["slide_css"] is None


@pytest.mark.parametrize(
    "bad",
    [
        ":host { position: fixed; inset: 0; z-index: 2147483647; }",
        ":/**/host { position: fixed; }",
        '.a::before{content:"/*"} :host{position:fixed;z-index:9} .b::before{content:"*/"}',
    ],
)
async def test_slide_css_may_not_style_the_host(client, tenant_id, bad):
    h = await _token(tenant_id, "approver")
    resp = await client.post("/v1/themes", json={**THEME, "slide_css": bad}, headers=h)
    assert resp.status_code == 422
    assert [p["message"] for p in resp.json()["problems"]] == [
        "slide CSS: :host is not allowed, nor a comment straight after ':'"
    ]


def test_the_host_check_is_linear():
    # Many comments after a colon with no host: a backtracking pattern takes minutes.
    started = time.perf_counter()
    assert check_slide_css(":" + "/**/" * 40 + "x")
    assert check_slide_css(":/*" * 25_000) and check_slide_css("a:hover{}" * 10_000) == []
    assert time.perf_counter() - started < 0.1


@pytest.mark.parametrize(
    "ok",
    [
        ".slide { background-image: linear-gradient(#fff, #eee); }",
        ".slide { list-style-image: none; border-image: none; mask-image: none; }",
        "@page { @top-center { content: element(docheader); } }",
    ],
)
async def test_ordinary_slide_css_still_passes(client, tenant_id, ok):
    h = await _token(tenant_id, "approver")
    resp = await client.post("/v1/themes", json={**THEME, "slide_css": ok}, headers=h)
    assert resp.status_code == 201, resp.text


@pytest.mark.parametrize("theme", BUILTIN_THEMES, ids=lambda t: t.key)
def test_every_builtin_theme_css_passes_the_check(theme):
    assert check_theme_css(theme.css or "") == []


async def test_a_builtin_theme_stays_editable(client, tenant_id):
    h = await _token(tenant_id, "approver")
    await client.get("/v1/themes", headers=h)
    resp = await client.patch("/v1/themes/roadmap-board", json={"description": "Ours"}, headers=h)
    assert resp.status_code == 200, resp.text


async def test_slide_css_is_stored_updated_cleared_and_audited(client, tenant_id):
    h = await _token(tenant_id, "approver")
    viewer = await _token(tenant_id, "viewer")
    resp = await client.post("/v1/themes", json={**THEME, "slide_css": SLIDE_CSS}, headers=h)
    assert resp.status_code == 201, resp.text
    assert resp.json()["slide_css"] == SLIDE_CSS
    got = (await client.get("/v1/themes/brand", headers=viewer)).json()
    assert got["slide_css"] == SLIDE_CSS
    assert {t["key"]: t["slide_css"] for t in (await client.get("/v1/themes", headers=h)).json()}[
        "clean-document"
    ] is None

    # Other edits keep it; a new value replaces it; null or blank clears it.
    resp = await client.patch("/v1/themes/brand", json={"css": "h2 { color: #222; }"}, headers=h)
    assert resp.json()["slide_css"] == SLIDE_CSS
    resp = await client.patch(
        "/v1/themes/brand", json={"slide_css": ".deck .slide { padding: 30px; }"}, headers=h
    )
    assert resp.status_code == 200 and resp.json()["slide_css"].startswith(".deck .slide {")
    resp = await client.patch("/v1/themes/brand", json={"slide_css": "  "}, headers=h)
    assert resp.json()["slide_css"] is None
    resp = await client.patch("/v1/themes/brand", json={"slide_css": SLIDE_CSS}, headers=h)
    resp = await client.patch("/v1/themes/brand", json={"slide_css": None}, headers=h)
    assert resp.json()["slide_css"] is None

    async with tenant_session(tenant_id) as s:
        events = list(
            await s.scalars(
                select(AuditEvent)
                .where(AuditEvent.action == "theme.updated")
                .order_by(AuditEvent.created_at)
            )
        )
    assert [e.data["fields"] for e in events] == [
        ["css"],
        ["slide_css"],
        ["slide_css"],
        ["slide_css"],
        ["slide_css"],
    ]


async def test_builtin_themes_keep_their_slide_variant(client, tenant_id):
    h = await _token(tenant_id, "approver")
    await seed_company(tenant_id)
    resp = await client.patch("/v1/themes/clean-document", json={"slide_css": SLIDE_CSS}, headers=h)
    assert resp.status_code == 422
    assert [p["code"] for p in resp.json()["problems"]] == ["builtin_slide_css"]
    # Other edits of a built-in still work, and clearing (a no-op) is accepted.
    resp = await client.patch(
        "/v1/themes/clean-document", json={"description": "Ours", "slide_css": None}, headers=h
    )
    assert resp.status_code == 200 and resp.json()["slide_css"] is None


async def test_theme_roles_are_unchanged(client, tenant_id):
    viewer = await _token(tenant_id, "viewer")
    commenter = await _token(tenant_id, "commenter")
    h = await _token(tenant_id, "approver")
    assert (await client.post("/v1/themes", json=THEME, headers=h)).status_code == 201
    for who in (viewer, commenter):
        payload = {**THEME, "key": "other", "slide_css": SLIDE_CSS}
        assert (await client.post("/v1/themes", json=payload, headers=who)).status_code == 403
        assert (
            await client.patch("/v1/themes/brand", json={"slide_css": SLIDE_CSS}, headers=who)
        ).status_code == 403
        assert (await client.get("/v1/themes/brand", headers=who)).status_code == 200
        preview = await client.post(
            "/v1/themes/preview",
            json={"tokens": {}, "slide_css": SLIDE_CSS, "view": "slides"},
            headers=who,
        )
        assert preview.status_code == 200
    anonymous = await client.patch("/v1/themes/brand", json={"slide_css": SLIDE_CSS})
    assert anonymous.status_code == 401


async def _themed_doc(client, tenant_id, h, slide_css: str | None) -> None:
    """The checkout draft in a company theme with ``slide_css``."""
    await seed_company(tenant_id)
    theme = {**THEME, "slide_css": slide_css}
    assert (await client.post("/v1/themes", json=theme, headers=h)).status_code == 201
    resp = await _import(client, h, DRAFT)
    assert resp.status_code == 201, resp.text
    async with tenant_session(tenant_id) as s:
        doc = await s.scalar(select(Document).where(Document.slug == "checkout"))
        doc.theme = "brand"
        await s.commit()


async def test_present_data_applies_slide_css_after_the_shared_layout(client, tenant_id):
    h = await _token(tenant_id, "approver")
    await _themed_doc(client, tenant_id, h, SLIDE_CSS)
    deck = (await client.get("/v1/documents/checkout/versions/1/deck", headers=h)).json()
    assert deck["theme"] == "brand" and deck["theme_class"] == "theme-brand"
    css = deck["css"]
    shared = css.index(".deck .slide {")
    tokens = css.index("--color-text: #111111")
    assert shared < tokens < css.index(SLIDE_CSS)
    assert css.rstrip().endswith(SLIDE_CSS)
    # A built-in's slide variant is not mixed in.
    assert ".theme-clean-document" not in css

    # Changing the theme's slide CSS changes Present mode at once; clearing drops it.
    await client.patch("/v1/themes/brand", json={"slide_css": None}, headers=h)
    deck = (await client.get("/v1/documents/checkout/versions/1/deck", headers=h)).json()
    assert "SLIDECSS" not in deck["css"]


async def test_stored_unsafe_slide_css_never_reaches_a_deck(client, tenant_id):
    h = await _token(tenant_id, "approver")
    await _themed_doc(client, tenant_id, h, SLIDE_CSS)
    async with tenant_session(tenant_id) as s:
        theme = await s.scalar(select(Theme).where(Theme.key == "brand"))
        theme.slide_css = "@import 'https://evil.example/x.css'; .slide { color: red; }"
        await s.commit()
    deck = (await client.get("/v1/documents/checkout/versions/1/deck", headers=h)).json()
    assert "@import" not in deck["css"]
    html = (
        await client.get(
            "/v1/documents/checkout/versions/1/export?layer=deck&format=html", headers=h
        )
    ).text
    assert "@import" not in html and "evil.example" not in html


async def test_deck_html_and_pptx_with_slide_css(client, tenant_id):
    h = await _token(tenant_id, "approver")
    await _themed_doc(client, tenant_id, h, SLIDE_CSS)
    base = "/v1/documents/checkout/versions/1/export"
    html = (await client.get(f"{base}?layer=deck&format=html", headers=h)).text
    assert SLIDE_CSS in html and 'class="deck deck-print theme-brand"' in html
    assert html.index(".deck .slide {") < html.index(SLIDE_CSS)
    # The PPTX keeps using the theme tokens; the slide CSS is for HTML slides only.
    pptx = await client.get(f"{base}?layer=deck&format=pptx", headers=h)
    prs = Presentation(io.BytesIO(pptx.content))
    texts = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    assert "SLIDECSS" not in texts and "Checkout PRD" in texts


@requires_pdf
async def test_deck_pdf_uses_the_slide_css(client, tenant_id):
    h = await _token(tenant_id, "approver")
    await _themed_doc(client, tenant_id, h, SLIDE_CSS)
    resp = await client.get(
        "/v1/documents/checkout/versions/1/export?layer=deck&format=pdf", headers=h
    )
    assert resp.status_code == 200 and resp.content.startswith(b"%PDF")
    pages = [p.extract_text() for p in PdfReader(io.BytesIO(resp.content)).pages]
    assert len(pages) == 1 + len(BRIEF_KEYS) + 1
    # The ::after rule of the slide CSS puts its marker after every slide title.
    assert "SLIDECSS" in pages[1] and "Summary" in pages[1]
    plain = await client.get(
        "/v1/documents/checkout/versions/1/export?layer=full&format=pdf", headers=h
    )
    assert "SLIDECSS" not in " ".join(
        p.extract_text() for p in PdfReader(io.BytesIO(plain.content)).pages
    )


async def test_slide_preview(client, tenant_id):
    h = await _token(tenant_id, "viewer")
    tokens = {"colors": {"accent": "#7a1fa2", "background": "#ffffff", "text": "#111111"}}
    resp = await client.post(
        "/v1/themes/preview",
        json={"tokens": tokens, "slide_css": SLIDE_CSS, "view": "slides"},
        headers=h,
    )
    assert resp.status_code == 200 and resp.headers["content-type"].startswith("text/html")
    html = resp.text
    assert SLIDE_CSS in html and "--color-accent: #7a1fa2" in html
    assert 'class="deck deck-preview theme-preview"' in html
    assert html.count('class="slide slide--section"') == 3  # the sample's brief
    assert "Saved carts PRD" in html
    roadmap = await client.post(
        "/v1/themes/preview",
        json={"tokens": {}, "slide_css": SLIDE_CSS, "view": "slides", "kind": "roadmap"},
        headers=h,
    )
    assert roadmap.text.count('class="slide slide--section"') == 3
    # Unsafe slide CSS is left out of the preview, like unsafe CSS in the document one.
    unsafe = await client.post(
        "/v1/themes/preview",
        json={"tokens": {}, "slide_css": "@import 'https://x.example/a.css';", "view": "slides"},
        headers=h,
    )
    assert unsafe.status_code == 200 and "@import" not in unsafe.text
    # A built-in theme's slides: its variant, not slide CSS.
    builtin = await client.post(
        "/v1/themes/preview",
        json={
            "tokens": {},
            "slide_css": SLIDE_CSS,
            "view": "slides",
            "slide_variant": "clean-document",
        },
        headers=h,
    )
    assert ".theme-clean-document .slide-title" in builtin.text and SLIDE_CSS not in builtin.text
    # The document preview is unchanged by default.
    doc = await client.post(
        "/v1/themes/preview", json={"tokens": {}, "slide_css": SLIDE_CSS}, headers=h
    )
    assert "slide--section" not in doc.text and SLIDE_CSS not in doc.text


def test_mcp_theme_tools_carry_slide_css():
    tools = {o.tool_name: o for o in app_operations(app)}
    for name in ("create_theme", "update_theme", "preview_theme"):
        schema = tools[name].input_schema()
        body = schema["properties"]["body"]
        ref = body["$ref"][len("#/$defs/") :]
        assert "slide_css" in schema["$defs"][ref]["properties"], name
    assert "view" in schema["$defs"][ref]["properties"]
    out = app.openapi()["paths"]["/v1/themes/{key}"]["get"]["responses"]["200"]
    ref = out["content"]["application/json"]["schema"]["$ref"].rsplit("/", 1)[1]
    assert ref == "ThemeOut"
    assert "slide_css" in app.openapi()["components"]["schemas"]["ThemeOut"]["properties"]
    assert "get_theme" in tools


# --- presenter notes only from notes blocks ---------------------------------------------


GUIDED = DRAFT.replace(
    "## Problem {#problem}\nReturning buyers",
    "## Problem {#problem}\n<!-- Say who has the problem, in plain language. -->\nReturning buyers",
)


async def test_notes_come_only_from_notes_blocks(client, tenant_id):
    h = await _token(tenant_id, "approver")
    await seed_company(tenant_id)
    assert (await _import(client, h, GUIDED)).status_code == 201
    v = (await client.get("/v1/documents/checkout/versions/1", headers=h)).json()
    notes = {s["section_key"]: s["notes"] for s in v["sections"]}
    assert notes["summary"] == "Open with the drop-off number."
    # Neither the draft's own guidance nor the template's is used as notes.
    assert notes["problem"] is None and notes["goal-trace"] is None
    assert all(notes[k] is None for k in notes if k not in BRIEF_KEYS)

    deck = (await client.get("/v1/documents/checkout/versions/1/deck", headers=h)).json()
    by_key = {s["section_key"]: s for s in deck["slides"] if s["section_key"]}
    assert by_key["summary"]["notes"] == "Open with the drop-off number."
    assert [k for k, s in by_key.items() if s["notes"]] == ["summary"]

    pptx = await client.get(
        "/v1/documents/checkout/versions/1/export?layer=deck&format=pptx", headers=h
    )
    slides = list(Presentation(io.BytesIO(pptx.content)).slides)
    assert slides[1].notes_slide.notes_text_frame.text == "Open with the drop-off number."
    # The problem slide (and every other one without a block) gets no speaker notes.
    for slide in slides[2:]:
        assert not slide.has_notes_slide or slide.notes_slide.notes_text_frame.text == ""


async def test_notes_rule_applies_to_versions_imported_before_o1g(client, tenant_id):
    h = await _token(tenant_id, "approver")
    await seed_company(tenant_id)
    assert (await _import(client, h, GUIDED)).status_code == 201
    # Notes are computed when a version is read, so an older version (whose template
    # has guidance for every section) shows none where it has no notes block.
    v = (await client.get("/v1/documents/checkout/versions/1", headers=h)).json()
    template = (await client.get("/v1/templates/oracle-default", headers=h)).json()
    assert "<!--" in template["versions"][-1]["body"]
    brief = [s for s in v["sections"] if s["layer"] == "brief"]
    assert [s["section_key"] for s in brief if s["notes"]] == ["summary"]


async def test_themeless_export_strips_notes_and_guidance(client, tenant_id):
    h = await _token(tenant_id, "approver")
    await seed_company(tenant_id)
    assert (await _import(client, h, GUIDED)).status_code == 201
    async with tenant_session(tenant_id) as s:
        ev = await load_export_version(s, "checkout", 1)
        ev.theme = None  # no theme resolves: the plain export
        html = render_html(ev)
    assert "One page for returning buyers" in html or "one page for returning buyers" in html
    assert "Open with the drop-off number" not in html
    assert "Say who has the problem" not in html and "&lt;!--" not in html
    assert "language-notes" not in html
