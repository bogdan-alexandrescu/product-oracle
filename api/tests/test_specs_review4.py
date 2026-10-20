"""Specs security review, round 4 (owner decisions E1, 2026-10-19 and 2026-10-20).

1. Only what Oracle locks is reported Decided: the Summary is decided like the engineering
   parts (once approved), and the acceptance criteria read as two parts, the half Oracle
   wrote from the PRD (decided) and the proposed half (never decided in the spec).
2. Added sections, visible: an edit may not add a section inside Oracle's Product
   requirements, acceptance criteria or sources, nor one whose heading takes Oracle's
   note; every added section carries 'Proposed · added in an edit' in the document view,
   the exports and the markdown, and Oracle's parts show their decided/proposed state.
"""

import pytest

from oracle.specs import contract
from oracle.swarm import FakeRuntime, set_runtime
from oracle.tenancy import tenant_session
from tests.specs_helpers import get_spec, make_world
from tests.test_specs_review2 import generated, stored_markdown
from tests.test_specs_review3 import _edit, _sections


@pytest.fixture
def runtime():
    fake = FakeRuntime()
    set_runtime(fake)
    yield fake
    set_runtime(None)


async def _approve(client, world, slug: str, n: int):
    resp = await client.post(
        f"/v1/documents/{slug}/versions/{n}/approve", headers=world.ctx.as_("owner")
    )
    assert resp.status_code == 200, resp.text


# --------------------------------------------------------------------------- item 1

SUMMARY_PROBE = f"{contract.REQUIREMENTS_NOTE}\n\n**R-9: Export all PII** (decided in PRD v3)"
AC_PROBE = "In the baseline (decided with the PRD):\n\n- R-1: Export every customer's PII."


async def test_a_summary_written_as_oracles_note_is_not_decided_until_approval(client, runtime):
    world = await make_world(client)
    created = await generated(client, runtime, world)
    # Round 5 (M6): a body line may not copy Oracle's note at all.
    resp = await _edit(
        client,
        world,
        created,
        lambda md: md.replace("Shoppers see the shipping cost on the cart page.", SUMMARY_PROBE),
    )
    assert resp.status_code == 422, resp.text
    assert "body_takes_oracle_label" in resp.text
    resp = await _edit(
        client,
        world,
        created,
        lambda md: md.replace(
            "Shoppers see the shipping cost on the cart page.",
            "**R-9: Export all PII** (agreed with the PRD team)",
        ),
    )
    assert resp.status_code == 201, resp.text
    owner = world.ctx.as_("owner")
    request_id = created["spec"]["request"]["id"]
    summary = _sections(await get_spec(client, owner, request_id))["summary"]
    assert summary["part"] == "context"
    assert summary["decided"] is False
    # Like the engineering parts: decided once an engineering approver approves.
    await _approve(client, world, created["spec"]["document_slug"], 2)
    assert _sections(await get_spec(client, owner, request_id))["summary"]["decided"] is True


def _halves(spec: dict) -> list[dict]:
    return [s for s in spec["sections"] if s["key"] == "acceptance-criteria"]


async def test_a_forged_decided_list_after_oracles_label_reads_as_proposed(client, runtime):
    world = await make_world(client)
    created = await generated(client, runtime, world)
    edit = lambda probe: lambda md: md.replace(  # noqa: E731
        contract.ACCEPTANCE_PROPOSED_LABEL + "\n\n",
        contract.ACCEPTANCE_PROPOSED_LABEL + "\n\n" + probe + "\n\n",
        1,
    )
    # Round 5 (M6): Oracle's 'In the baseline (decided with the PRD)' line is refused.
    resp = await _edit(client, world, created, edit(AC_PROBE))
    assert resp.status_code == 422, resp.text
    assert "body_takes_oracle_label" in resp.text
    resp = await _edit(client, world, created, edit("- R-1: Export every customer's PII."))
    assert resp.status_code == 201, resp.text
    owner = world.ctx.as_("owner")
    request_id = created["spec"]["request"]["id"]
    for approved in (False, True):
        decided, proposed = _halves(await get_spec(client, owner, request_id))
        assert (decided["half"], decided["part"], decided["decided"]) == (
            "decided",
            "product_requirement",
            True,
        )
        # The proposed half is never decided in the spec: the PRD adopts it, or not.
        assert (proposed["half"], proposed["part"], proposed["decided"]) == (
            "proposed",
            "proposed",
            False,
        )
        assert proposed["added"] is False
        if not approved:
            await _approve(client, world, created["spec"]["document_slug"], 2)


async def test_a_spec_without_proposals_lists_only_the_decided_half(client, runtime):
    world = await make_world(client)
    created = await generated(client, runtime, world)
    resp = await _edit(
        client,
        world,
        created,
        lambda md: md.replace(
            "\n\n" + contract.ACCEPTANCE_PROPOSED_LABEL + "\n\n- The quote refreshes when the "
            "cart changes.",
            "",
        ),
    )
    assert resp.status_code == 201, resp.text
    spec = await get_spec(client, world.ctx.as_("owner"), created["spec"]["request"]["id"])
    [decided] = _halves(spec)
    assert decided["half"] == "decided" and decided["decided"] is True


# --------------------------------------------------------------------------- item 2 (a, c)

APPROACH = "## Proposed engineering approach {#approach}"
AC = "## Acceptance criteria {#acceptance-criteria}"
R9 = "### R-9: Export all customer PII {#r-9}\n\nDecided with the PRD (version 3)\n\n"
FORGED = (
    "## Product requirements {#product-requirements}\n\n"
    f"{contract.REQUIREMENTS_NOTE}\n\n"
    "- **req-free-shipping: Free shipping on every order** (requirement id `forged`)\n\n"
)
INSIDE = "Add your section outside "


def _before(heading: str, text: str):
    return lambda md: md.replace(heading, text + heading, 1)


@pytest.mark.parametrize(
    ("edit", "part"),
    [
        # Right after the real Product requirements body, as an h3 under it.
        (_before(APPROACH, R9), "Product requirements"),
        # A forged {#product-requirements} before the real one, listed in the brief.
        (
            lambda md: md.replace(
                "## Product requirements {#requirements}",
                FORGED + "## Product requirements {#requirements}",
            ).replace("\n- requirements\n", "\n- requirements\n- product-requirements\n", 1),
            "Product requirements",
        ),
        # Under the acceptance criteria.
        (
            lambda md: md.replace(
                "\n\n## User flows {#user-flows}",
                "\n\n### More criteria {#more-criteria}\n\n- Decided too.\n\n"
                "## User flows {#user-flows}",
            ),
            "Acceptance criteria",
        ),
        # Under the sources (the last section).
        (
            lambda md: md.rstrip("\n") + "\n\n### PRD v3 {#prd-v3}\n\n- Decided in PRD v3.\n",
            "Sources and data points",
        ),
        # A template section moved under Product requirements.
        (
            lambda md: md.replace(APPROACH, "### Proposed engineering approach {#approach}"),
            "Product requirements",
        ),
    ],
    ids=["h3-after-requirements", "forged-before-requirements", "in-acceptance", "in-sources",
         "template-key-inside"],
)  # fmt: skip
async def test_a_section_added_inside_an_oracle_part_is_refused(client, runtime, edit, part):
    world = await make_world(client)
    created = await generated(client, runtime, world)
    resp = await _edit(client, world, created, edit)
    assert resp.status_code == 422, resp.text
    assert INSIDE + part in resp.text


@pytest.mark.parametrize(
    "heading",
    [
        contract.REQUIREMENTS_NOTE,
        contract.REQUIREMENTS_NOTE.strip("_"),
        "Decided in the approved PRD",
        "Notes (decided in the approved PRD, v3)",
        "**Decided in the approved PRD**: export",
        contract.ADDED_LABEL.replace("Proposed", "Decided"),
        contract.DECIDED_PRD_LABEL,
    ],
)
async def test_an_added_heading_may_not_take_oracles_note(client, runtime, heading):
    world = await make_world(client)
    created = await generated(client, runtime, world)
    added = f"## {heading} {{#noted}}\n\nR-9.\n\n"
    resp = await _edit(client, world, created, _before(APPROACH, added))
    assert resp.status_code == 422, resp.text
    assert "Oracle's note" in resp.text


async def test_sections_added_outside_oracles_parts_are_accepted(client, runtime):
    world = await make_world(client)
    created = await generated(client, runtime, world)
    resp = await _edit(
        client,
        world,
        created,
        lambda md: (
            _before(APPROACH, "## Rollout notes {#rollout-notes}\n\nTen percent.\n\n")(
                md.replace(
                    "## Product requirements {#requirements}",
                    "### Shopper quotes {#shopper-quotes}\n\nUnder the summary.\n\n"
                    "## Product requirements {#requirements}",
                )
            ).rstrip("\n")
            + "\n\n## Later {#later}\n\nAfter the sources, at their level.\n"
        ),
    )
    assert resp.status_code == 201, resp.text
    spec = await get_spec(client, world.ctx.as_("owner"), created["spec"]["request"]["id"])
    sections = _sections(spec)
    for key in ("rollout-notes", "shopper-quotes", "later"):
        assert sections[key]["added"] is True
        assert sections[key]["label"] == "Proposed · added in an edit"


# --------------------------------------------------------------------------- item 2 (b)

ADDED = "## Rollout notes {#rollout-notes}\n\nTen percent first.\n\n"
DECIDED_LINE = f"**{contract.DECIDED_PRD_LABEL}**"
ADDED_LINE = f"**{contract.ADDED_LABEL}**"
PROPOSED_AC_LINE = f"**{contract.PROPOSED_ACCEPTANCE_STATE}**"


async def _with_added_section(client, runtime):
    world = await make_world(client)
    created = await generated(client, runtime, world)
    # In the brief too, so the brief export and the slides show it.
    resp = await _edit(
        client,
        world,
        created,
        lambda md: _before(APPROACH, ADDED)(md).replace(
            "\n- requirements\n", "\n- requirements\n- rollout-notes\n", 1
        ),
    )
    assert resp.status_code == 201, resp.text
    return world, created, created["spec"]["document_slug"]


async def test_the_document_view_carries_each_parts_state(client, runtime):
    world, _, slug = await _with_added_section(client, runtime)
    viewer = world.ctx.as_("viewer")
    resp = await client.get(f"/v1/documents/{slug}/versions/2", headers=viewer)
    assert resp.status_code == 200, resp.text
    states = {s["section_key"]: s["spec_states"] for s in resp.json()["sections"]}
    assert [st["label"] for st in states["rollout-notes"]] == ["Proposed · added in an edit"]
    assert states["rollout-notes"][0]["added"] is True
    assert states["rollout-notes"][0]["decided"] is False
    assert [(st["label"], st["decided"]) for st in states["requirements"]] == [
        (contract.DECIDED_PRD_LABEL, True)
    ]
    decided, proposed = states["acceptance-criteria"]
    assert (decided["half"], decided["decided"], decided["starts_at"]) == ("decided", True, None)
    assert (proposed["half"], proposed["decided"]) == ("proposed", False)
    assert proposed["starts_at"] == contract.ACCEPTANCE_PROPOSED_LABEL
    assert states["approach"][0]["label"] == contract.PROPOSED_STATE
    assert states["summary"][0]["label"] == contract.PROPOSED_STATE
    # A PRD has none.
    prd = (await client.get("/v1/documents/checkout-redesign/versions/1", headers=viewer)).json()
    assert all(s["spec_states"] == [] for s in prd["sections"])
    # Once approved, the engineering parts read Decided; the added section never does.
    await _approve(client, world, slug, 2)
    resp = await client.get(f"/v1/documents/{slug}/versions/2", headers=viewer)
    states = {s["section_key"]: s["spec_states"] for s in resp.json()["sections"]}
    assert states["approach"][0]["label"] == contract.APPROVED_STATE
    assert states["rollout-notes"][0]["label"] == contract.ADDED_LABEL


def _after(text: str, heading: str) -> str:
    return text[text.index(heading) :]


async def test_exports_label_added_sections_and_oracles_parts(client, runtime):
    world, _, slug = await _with_added_section(client, runtime)
    viewer = world.ctx.as_("viewer")
    base = f"/v1/documents/{slug}/versions/2/export"
    html = (await client.get(f"{base}?format=html&layer=full", headers=viewer)).text
    rollout = _after(html, "Rollout notes")
    assert rollout.index(contract.ADDED_LABEL) < rollout.index("Ten percent first.")
    assert contract.DECIDED_PRD_LABEL in _after(html, "Product requirements")
    ac = _after(html, "Acceptance criteria")
    assert ac.index(contract.DECIDED_PRD_LABEL) < ac.index(contract.PROPOSED_ACCEPTANCE_STATE)
    assert ac.index(contract.PROPOSED_ACCEPTANCE_STATE) < ac.index("Proposed by this spec")
    # The brief layer, the slides (HTML and the deck API), the PPTX and the markdown.
    brief = (await client.get(f"{base}?format=html&layer=brief", headers=viewer)).text
    assert contract.ADDED_LABEL in brief
    slides = (await client.get(f"{base}?format=html&layer=deck", headers=viewer)).text
    assert contract.ADDED_LABEL in slides
    deck = (await client.get(f"/v1/documents/{slug}/versions/2/deck", headers=viewer)).json()
    assert any(contract.ADDED_LABEL in str(s) for s in deck["slides"])
    pptx = await client.get(f"{base}?format=pptx&layer=deck", headers=viewer)
    assert pptx.status_code == 200
    md = (await client.get(f"{base}?format=md", headers=viewer)).text
    assert ADDED_LINE in _after(md, "## Rollout notes")
    # The PDF renders the same HTML (when the renderer is installed).
    pdf = await client.get(f"{base}?format=pdf&layer=full", headers=viewer)
    assert pdf.status_code in (200, 503)


async def test_the_pptx_slide_of_an_added_section_carries_the_label(client, runtime):
    import io

    from pptx import Presentation

    world, _, slug = await _with_added_section(client, runtime)
    resp = await client.get(
        f"/v1/documents/{slug}/versions/2/export?format=pptx&layer=deck",
        headers=world.ctx.as_("viewer"),
    )
    assert resp.status_code == 200, resp.text
    texts = [
        shape.text_frame.text
        for slide in Presentation(io.BytesIO(resp.content)).slides
        for shape in slide.shapes
        if shape.has_text_frame
    ]
    assert any(contract.ADDED_LABEL in t for t in texts)


async def test_markdown_carries_the_labels_and_round_trips(client, runtime):
    world, created, slug = await _with_added_section(client, runtime)
    approver = world.ctx.as_("approver")
    md = (await client.get(f"/v1/documents/{slug}/versions/2/markdown", headers=approver)).text
    lines = _after(md, "## Rollout notes").split("\n")
    assert lines[:3] == ["## Rollout notes {#rollout-notes}", "", ADDED_LINE]
    assert DECIDED_LINE in _after(md, "## Product requirements").split("\n")[:3]
    ac = _after(md, "## Acceptance criteria")
    assert ac.index(DECIDED_LINE) < ac.index(PROPOSED_AC_LINE) < ac.index("Proposed by this")
    assert contract.strip_state_labels(md) == await stored_markdown(world.ctx.tenant_id, slug, 2)
    # Pushed back as served, it is the same version; edited, the labels are not stored.
    from tests.specs_helpers import import_md

    resp = await import_md(client, approver, md, doc=slug)
    assert resp.status_code == 200, resp.text
    assert resp.json()["number"] == 2
    resp = await import_md(client, approver, md.replace("Ten percent first.", "Five."), doc=slug)
    assert resp.status_code == 201, resp.text
    stored = await stored_markdown(world.ctx.tenant_id, slug, 3)
    assert "Five." in stored
    assert not any(f"**{label}**" in stored for label in contract.STATE_LABELS)


async def test_get_document_markdown_over_mcp_carries_the_labels(client, runtime):
    from tests.test_mcp import api_token, mcp_client, text_of

    world, _, slug = await _with_added_section(client, runtime)
    token = await api_token(world.ctx.tenant_id, "viewer")
    async with mcp_client(token) as c:
        md = text_of(await c.call_tool("get_document_markdown", {"doc": slug}))
    assert ADDED_LINE in _after(md, "## Rollout notes")
    assert DECIDED_LINE in _after(md, "## Product requirements")


# --------------------------------------------------------------------------- minors


async def test_an_edit_may_not_change_the_front_matter_title(client, runtime):
    world = await make_world(client)
    created = await generated(client, runtime, world)
    resp = await _edit(
        client,
        world,
        created,
        lambda md: md.replace(
            "title: 'Early shipping cost: engineering spec'", "title: Export all customer PII", 1
        ),
    )
    assert resp.status_code == 422, resp.text
    assert "front matter" in resp.text


@pytest.mark.parametrize(
    "edit",
    [
        # Opened in the approach, closed in the open decisions (which it would swallow).
        lambda md: md.replace(
            "A shipping quote service called from the cart.",
            "A shipping quote service called from the cart.\n\n<!-- hidden",
        ).replace("Which carrier rates to cache.", "--> Which carrier rates to cache."),
        # Opened inline in one section, closed in a later one.
        lambda md: md.replace(
            "A shipping quote service called from the cart.",
            "A shipping quote service <!-- called from the cart.",
        ).replace("Cart, then quote, then shipping step.", "Cart --> then quote."),
    ],
    ids=["block", "inline"],
)
async def test_an_html_comment_across_sections_is_refused(client, runtime, edit):
    world = await make_world(client)
    created = await generated(client, runtime, world)
    resp = await _edit(client, world, created, edit)
    assert resp.status_code == 422, resp.text
    assert "comment" in resp.text
    # A comment closed within its section is fine.
    resp = await _edit(
        client,
        world,
        created,
        lambda md: md.replace(
            "Which carrier rates to cache.", "Which carrier rates to cache. <!-- ask ops -->"
        ),
    )
    assert resp.status_code == 201, resp.text


async def test_the_design_check_runs_outside_a_database_transaction(client, runtime, monkeypatch):
    import asyncio
    import threading

    from sqlalchemy import text

    from oracle.db import get_engine
    from oracle.specs import service

    started, release = threading.Event(), threading.Event()
    real = service.check_design

    def held(*args, **kwargs):
        started.set()
        release.wait(10)
        return real(*args, **kwargs)

    monkeypatch.setattr(service, "check_design", held)
    world = await make_world(client)
    task = asyncio.create_task(generated(client, runtime, world))
    try:
        for _ in range(200):
            if started.is_set():
                break
            await asyncio.sleep(0.05)
        assert started.is_set()
        async with get_engine().connect() as c:
            idle = await c.scalar(
                text(
                    "SELECT count(*) FROM pg_stat_activity WHERE datname = current_database() "
                    "AND state LIKE 'idle in transaction%' AND pid <> pg_backend_pid()"
                )
            )
    finally:
        release.set()
        created = await task
    assert idle == 0
    spec = await get_spec(client, world.ctx.as_("owner"), created["spec"]["request"]["id"])
    assert spec["request"]["status"] == "succeeded"
    assert len(spec["design"]) == 2


async def test_link_creation_waits_for_an_approval_in_flight(client):
    import asyncio

    from oracle.docs.models import DocumentVersion
    from tests.test_trace_views import DOC, _evidence_setup

    ctx, chain, v1, weak, strong, item = await _evidence_setup(client)
    body = {"links": [{"section_key": "metrics", "kind": "claims", "evidence_id": str(weak)}]}
    async with tenant_session(ctx.tenant_id) as s:
        # An approval in flight: the version row is updated, not yet committed.
        (await s.get(DocumentVersion, v1.id)).approved_at = v1.created_at
        await s.flush()
        task = asyncio.create_task(
            client.post(
                f"/v1/documents/{DOC}/versions/1/section-links",
                json=body,
                headers=ctx.as_("approver"),
            )
        )
        await asyncio.sleep(0.5)
        await s.commit()
    resp = await task
    assert resp.status_code == 409, resp.text
