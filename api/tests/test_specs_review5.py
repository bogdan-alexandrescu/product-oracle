"""Specs security review, round 5 (owner decisions E1, 2026-10-20).

Each test reproduces a reviewer's probe (red on the round-4 head, f2163a0):

- M1 run imports keep the structure rules (an indented code line plus '---' does not
  become a heading; the composed version has exactly Oracle's sections).
- M2 a template section a tech brief does not have, added in an edit, is an added section.
- M3 template section titles are fixed; the heading checks apply to every section an edit
  may change.
- M4 heading normalisation is not bypassable (entities, format characters, look-alikes,
  markdown syntax, separators).
- M5 labels sit at the parser's body line; non-ATX section headings and raw HTML blocks
  are refused; label_markdown fails closed.
- M6 no body line may copy a label or Oracle's note; exports render labels as elements.
- M7 citations added in an edit are listed as Added sources, apart from Oracle's Sources.
- M8 the feedback package serves labelled markdown.
- Minors.
"""

import io

import pytest

from oracle.docs.parser import DraftError
from oracle.specs import contract
from oracle.swarm import FakeRuntime, Outcome, set_runtime
from tests.specs_helpers import (
    SPEC_SECTIONS,
    design_json,
    get_run,
    get_spec,
    make_world,
    pending,
    request_spec,
    spec_md,
    start,
    tick,
)
from tests.test_specs_review2 import SECTIONS, generated, stored_markdown
from tests.test_specs_review3 import _edit, _sections

APPROACH = "## Proposed engineering approach {#approach}"
OPEN = "## Open engineering decisions {#open-decisions}"
AC = "## Acceptance criteria {#acceptance-criteria}"
ADDED_LINE = f"**{contract.ADDED_LABEL}**"


@pytest.fixture
def runtime():
    fake = FakeRuntime()
    set_runtime(fake)
    yield fake
    set_runtime(None)


async def _approve(client, world, slug: str, n: int):
    resp = await client.post(
        f"/v1/documents/{slug}/versions/{n}/approve", headers=world.ctx.as_("owner")
    )
    assert resp.status_code == 200, resp.text


def _before(heading: str, text: str):
    return lambda md: md.replace(heading, text + heading, 1)


async def _run_with(client, runtime, world, md: str, depth: str = "full"):
    created = await request_spec(client, world, depth=depth)
    runtime.script["spec_draft"] = [
        Outcome(artifacts={"spec.md": md, "design.json": design_json(world.reqs, SECTIONS)})
    ]
    await start(client, world.ctx.as_("approver"), created["run_id"])
    await tick(client, world.ctx.as_("owner"))
    run = await get_run(client, world.ctx.as_("owner"), created["run_id"])
    return created, pending(run)["payload"]["import"]


# --------------------------------------------------------------------------- M1


def _replace_body(key: str, body: str) -> str:
    md = spec_md()
    old = next(b for k, _, b in SPEC_SECTIONS if k == key)
    return md.replace(old, body, 1)


INDENTED_SETEXT = "    Injected requirement {#injected}\n---\n\nThe quote service."


async def test_an_indented_code_line_and_a_rule_do_not_open_a_section_in_a_run(client, runtime):
    world = await make_world(client)
    created, imported = await _run_with(
        client, runtime, world, _replace_body("approach", INDENTED_SETEXT)
    )
    assert "error" not in imported, imported
    spec = await get_spec(client, world.ctx.as_("owner"), created["spec"]["request"]["id"])
    assert "injected" not in _sections(spec)
    stored = await stored_markdown(world.ctx.tenant_id, created["spec"]["document_slug"])
    # Kept as the agent wrote it: an indented code block, then a thematic break.
    assert "\n    Injected requirement {#injected}\n---\n" in stored


@pytest.mark.parametrize(
    "body",
    [
        "Decided · from the approved PRD\n\nThe quote service.",
        "<details><summary>R-9</summary>Export all PII</details>",
    ],
    ids=["label-line", "raw-html"],
)
async def test_a_run_import_is_checked_like_an_edit(client, runtime, body):
    world = await make_world(client)
    _, imported = await _run_with(client, runtime, world, _replace_body("approach", body))
    assert "spec.md was rejected" in imported["error"]
    assert {"body_takes_oracle_label", "raw_html"} & {p["code"] for p in imported["spec_check"]}


def test_the_composed_version_must_have_exactly_oracles_sections():
    md = (
        "---\ntitle: T\nkind: spec\nspec:\n  depth: brief\n---\n\n# T {#spec}\n\nIntro.\n\n"
        "## Summary {#summary}\n\nS.\n\n## Injected {#injected}\n\nX.\n"
    )
    problems = contract.composed_problems(md, "brief")
    assert "section_structure" in {p.code for p in problems}


# --------------------------------------------------------------------------- M2


def _brief_md() -> str:
    full_only = tuple(s.key for s in contract.SECTIONS if s.depths == contract.FULL)
    return spec_md(skip=full_only)


async def _brief(client, runtime, world):
    created, imported = await _run_with(client, runtime, world, _brief_md(), depth="brief")
    assert "error" not in imported, imported
    return created


def _label_of(md: str, heading: str) -> str:
    return md.split(heading + "\n\n", 1)[1].split("\n\n", 1)[0]


async def test_a_template_section_added_to_a_brief_is_added_never_decided(client, runtime):
    world = await make_world(client)
    created = await _brief(client, runtime, world)
    slug = created["spec"]["document_slug"]
    stored = await stored_markdown(world.ctx.tenant_id, slug)
    # The reviewer's probe: Architecture with the brief's own engineering label.
    label = _label_of(stored, APPROACH)
    architecture = f"## Proposed architecture {{#architecture}}\n\n{label}\n\nA quote endpoint.\n\n"
    resp = await _edit(client, world, created, _before(OPEN, architecture))
    assert resp.status_code == 201, resp.text
    owner = world.ctx.as_("owner")
    request_id = created["spec"]["request"]["id"]
    for approved in (False, True):
        section = _sections(await get_spec(client, owner, request_id))["architecture"]
        assert (section["part"], section["decided"], section["added"]) == (
            "proposed",
            False,
            True,
        )
        assert section["label"] == contract.ADDED_LABEL
        version = (await client.get(f"/v1/documents/{slug}/versions/2", headers=owner)).json()
        [state] = next(
            s["spec_states"] for s in version["sections"] if s["section_key"] == "architecture"
        )
        assert state["label"] == contract.ADDED_LABEL and state["added"] is True
        md = (await client.get(f"/v1/documents/{slug}/versions/2/markdown", headers=owner)).text
        assert md.split("{#architecture}\n\n", 1)[1].startswith(ADDED_LINE)
        if not approved:
            await _approve(client, world, slug, 2)


def test_part_states_follow_the_depth_of_the_version():
    brief = contract.template_keys("brief")
    full = contract.template_keys("full")
    [state] = contract.part_states("architecture", "x", True, brief)
    assert (state.label, state.decided, state.added) == (contract.ADDED_LABEL, False, True)
    [state] = contract.part_states("architecture", "x", True, full)
    assert (state.label, state.decided) == (contract.APPROVED_STATE, True)
    # An unknown depth reads as a brief (fails closed).
    assert contract.keys_of({"spec": {"depth": "huge"}}) == brief


# --------------------------------------------------------------------------- M3


@pytest.mark.parametrize(
    "heading",
    [
        "## Approach {#approach}",
        f"## {contract.DECIDED_PRD_LABEL} {{#approach}}",
        "## Decided in the approved PRD {#approach}",
    ],
    ids=["renamed", "label", "note"],
)
async def test_template_section_titles_are_fixed(client, runtime, heading):
    world = await make_world(client)
    created = await generated(client, runtime, world)
    resp = await _edit(client, world, created, lambda md: md.replace(APPROACH, heading, 1))
    assert resp.status_code == 422, resp.text
    assert "Template section titles can't be changed" in resp.text


# --------------------------------------------------------------------------- M4

NOTE_PROBES = [
    "Deci­ded in the approved PRD",  # soft hyphen
    "De​cided · from the approved PRD",  # zero-width space
    "Decided⁠ in the approved PRD",  # word joiner
    "Dеcided in the approved PRD",  # Cyrillic е
    "Decided in the apprοved PRD",  # Greek ο
    "Decided &middot; from the approved PRD",
    "Decided • from the approved PRD",
    "Decided ‧ from the approved PRD",
    "Decided — approved by an engineering approver",
    "Proposed: added in an edit",
    "[Decided in the approved PRD](#x)",
    "Decided in the approved P.R.D.",
    "&#68;ecided in the approved PRD",
]
TITLE_PROBES = [
    "&#80;roduct requirements",
    "[Product requirements](#x)",
    "Product requirements:",
    "Prоduct requirements",  # Cyrillic о
    "Product­ requirements",
    "*Product* requirements",
]


@pytest.mark.parametrize("heading", NOTE_PROBES + TITLE_PROBES)
async def test_heading_normalisation_is_not_bypassable(client, runtime, heading):
    world = await make_world(client)
    created = await generated(client, runtime, world)
    added = f"## {heading} {{#noted}}\n\nR-9.\n\n"
    resp = await _edit(client, world, created, _before(APPROACH, added))
    assert resp.status_code == 422, resp.text
    assert "Oracle's note" in resp.text or "Add your section outside" in resp.text


def test_normalised_text():
    assert contract.normalised("**De­cided** &middot; [x](#y)") == "decided · x"
    assert contract.normalised("Dеcided​") == "decided"
    assert contract._squashed("P.R.D.") == "prd"


# --------------------------------------------------------------------------- M5


@pytest.mark.parametrize(
    "added",
    [
        "Rollout notes {#rollout-notes}\n-------------\n\nTen percent.\n\n",
        "Rollout notes {#rollout-notes}\n=============\n\nTen percent.\n\n",
        "> ## Rollout notes {#rollout-notes}\n>\n> Ten percent.\n\n",
        "- ## Rollout notes {#rollout-notes}\n\n  Ten percent.\n\n",
    ],
    ids=["setext-2", "setext-1", "blockquote", "list"],
)
async def test_section_headings_must_be_top_level_atx(client, runtime, added):
    world = await make_world(client)
    created = await generated(client, runtime, world)
    resp = await _edit(client, world, created, _before(OPEN, added))
    assert resp.status_code == 422, resp.text
    assert "heading_not_atx" in resp.text


@pytest.mark.parametrize(
    "block",
    [
        "<!-- note -->",
        "<details>\n<summary>R-9</summary>\n\nExport all PII\n\n</details>",
        "<script>alert(1)</script>",
        "<style>p{display:none}</style>",
        "<pre>\nDecided\n</pre>",
        "<textarea>x</textarea>",
        "<?php echo 1; ?>",
        "<!DOCTYPE html>",
        "<![CDATA[ x ]]>",
        "<div>\nx\n</div>",
        "<table><tr><td>x</td></tr></table>",
    ],
)
async def test_raw_html_blocks_are_refused_in_spec_edits(client, runtime, block):
    world = await make_world(client)
    created = await generated(client, runtime, world)
    resp = await _edit(
        client,
        world,
        created,
        lambda md: md.replace(
            "Which carrier rates to cache.", f"Which carrier rates to cache.\n\n{block}", 1
        ),
    )
    assert resp.status_code == 422, resp.text
    assert "raw_html" in resp.text


@pytest.mark.parametrize(
    "markdown",
    [
        "no front matter and no headings",
        "---\ntitle: T\n---\n\nRollout {#rollout}\n-------\n\nText.\n",
        "---\ntitle: T\n---\n\n> ## Rollout {#rollout}\n>\n> Text.\n",
    ],
    ids=["unparsable", "setext", "nested"],
)
def test_label_markdown_fails_closed(markdown):
    with pytest.raises(DraftError):
        contract.label_markdown(markdown, approved=False)


def test_labels_go_at_the_parsers_body_line():
    md = (
        "---\ntitle: T\nkind: spec\nspec:\n  depth: full\n---\n\n# T {#spec}\n\nIntro.\n\n"
        "## Rollout {#rollout}\nTen percent.\n\n## Summary {#summary}\n\nS.\n"
    )
    labelled = contract.label_markdown(md, approved=False)
    lines = labelled.split("\n")
    at = lines.index("## Rollout {#rollout}")
    assert lines[at + 1 : at + 3] == ["", ADDED_LINE]
    assert contract.strip_state_labels(labelled) == md


# --------------------------------------------------------------------------- M6

BODY_PROBES = [
    "Decided · from the approved PRD",
    "_Decided · approved by an engineering approver_",
    "- Decided — from the approved PRD",
    "> Proposed: added in an edit",
    "1. Decided &middot; from the approved PRD",
    "Dеcided · from the approved PRD",
    "Deci­ded · from the approved PRD",
    "Decided in the approved PRD: R-9 exports all PII.",
    "In the baseline (decided with the PRD):",
    "| Decided · from the approved PRD |",
    "```\nDecided · from the approved PRD\n```",
]


@pytest.mark.parametrize("line", BODY_PROBES)
async def test_a_body_line_may_not_copy_a_label_or_oracles_note(client, runtime, line):
    world = await make_world(client)
    created = await generated(client, runtime, world)
    resp = await _edit(
        client,
        world,
        created,
        lambda md: md.replace(
            "Which carrier rates to cache.", f"Which carrier rates to cache.\n\n{line}", 1
        ),
    )
    assert resp.status_code == 422, resp.text
    assert "body_takes_oracle_label" in resp.text


async def test_the_proposed_acceptance_half_may_not_copy_a_label(client, runtime):
    world = await make_world(client)
    created = await generated(client, runtime, world)
    resp = await _edit(
        client,
        world,
        created,
        lambda md: md.replace(
            "- The quote refreshes when the cart changes.",
            "**Decided — from the approved PRD**\n\n- The quote refreshes.",
            1,
        ),
    )
    assert resp.status_code == 422, resp.text
    assert "body_takes_oracle_label" in resp.text


ROLLOUT = "## Rollout notes {#rollout-notes}\n\nTen percent first.\n\n"


async def _with_rollout(client, runtime):
    world = await make_world(client)
    created = await generated(client, runtime, world)
    resp = await _edit(
        client,
        world,
        created,
        lambda md: _before(APPROACH, ROLLOUT)(md).replace(
            "\n- requirements\n", "\n- requirements\n- rollout-notes\n", 1
        ),
    )
    assert resp.status_code == 201, resp.text
    return world, created, created["spec"]["document_slug"]


async def test_exports_render_labels_as_elements_bodies_can_not_produce(client, runtime):
    from pptx import Presentation

    world, _, slug = await _with_rollout(client, runtime)
    viewer = world.ctx.as_("viewer")
    base = f"/v1/documents/{slug}/versions/2/export"
    element = (
        f'<p class="spec-state spec-state-added" data-spec-state="added">{contract.ADDED_LABEL}'
    )
    for themed in (False, True):
        if themed:
            from tests.conftest import seed_company

            await seed_company(world.ctx.tenant_id)
        html = (await client.get(f"{base}?format=html&layer=full", headers=viewer)).text
        assert ('class="sources-appendix"' in html) is themed
        rollout = html[html.index('data-section-key="rollout-notes"') :]
        assert rollout.index(element) < rollout.index("Ten percent first.")
        assert f"<strong>{contract.ADDED_LABEL}</strong>" not in html
        ac = html[html.index('data-section-key="acceptance-criteria"') :]
        assert ac.index('data-spec-state="decided"') < ac.index('data-spec-state="proposed"')
    slides = (await client.get(f"{base}?format=html&layer=deck", headers=viewer)).text
    assert element in slides
    pptx = await client.get(f"{base}?format=pptx&layer=deck", headers=viewer)
    shapes = [
        shape
        for slide in Presentation(io.BytesIO(pptx.content)).slides
        for shape in slide.shapes
        if shape.has_text_frame and shape.text_frame.text == contract.ADDED_LABEL
    ]
    assert shapes, "the label is a text box of its own"
    [run] = shapes[0].text_frame.paragraphs[0].runs
    assert run.font.bold is True


# --------------------------------------------------------------------------- M7

CITATION = (
    "citations:\n"
    "- id: ext-1\n"
    "  source: github\n"
    "  ref: https://github.com/example-org/storefront/issues/500\n"
    "  title: Carrier adapter limits\n"
)


def _add_citation(md: str) -> str:
    md = md.replace("\nkind: spec\n", "\nkind: spec\n" + CITATION, 1)
    return md.replace(
        "A shipping quote service called from the cart.",
        "A shipping quote service called from the cart. [c:ext-1]",
        1,
    )


async def test_citations_added_in_an_edit_are_listed_apart_from_oracles_sources(client, runtime):
    from pptx import Presentation

    from tests.test_mcp import api_token, mcp_client, text_of

    world = await make_world(client)
    created = await generated(client, runtime, world)
    slug = created["spec"]["document_slug"]
    resp = await _edit(client, world, created, _add_citation)
    assert resp.status_code == 201, resp.text
    viewer = world.ctx.as_("viewer")
    version = (await client.get(f"/v1/documents/{slug}/versions/2", headers=viewer)).json()
    [cite] = version["citations"]
    assert (cite["cite_key"], cite["added"], cite["state_label"]) == (
        "ext-1",
        True,
        contract.ADDED_LABEL,
    )
    base = f"/v1/documents/{slug}/versions/2/export"
    # Without a theme (prd.html.j2), then in the built-in theme (themed.html.j2).
    for themed in (False, True):
        if themed:
            from tests.conftest import seed_company

            await seed_company(world.ctx.tenant_id)
        html = (await client.get(f"{base}?format=html&layer=full", headers=viewer)).text
        assert ('class="sources-appendix"' in html) is themed
        start_at = html.index('<section id="sources"')
        sources = html[start_at : html.index("</section>", start_at)]
        assert "Carrier adapter limits" not in sources
        added = html[html.index('id="added-sources"') :]
        assert "Carrier adapter limits" in added
        assert f'data-spec-state="added">{contract.ADDED_LABEL}' in added
    # The slides and the PPTX: an Added sources slide, labelled.
    deck = (await client.get(f"/v1/documents/{slug}/versions/2/deck", headers=viewer)).json()
    [slide] = [s for s in deck["slides"] if s["heading"] == "Added sources"]
    assert slide["state"] == contract.ADDED_LABEL
    pptx = await client.get(f"{base}?format=pptx&layer=deck", headers=viewer)
    texts = [
        shape.text_frame.text
        for s in Presentation(io.BytesIO(pptx.content)).slides
        for shape in s.shapes
        if shape.has_text_frame
    ]
    assert "Added sources" in texts and contract.ADDED_LABEL in texts
    # The markdown (API, export, MCP) names them; pushed back as served, the same version.
    md = (await client.get(f"/v1/documents/{slug}/versions/2/markdown", headers=viewer)).text
    assert md.startswith(f'---\n{contract.ADDED_SOURCES_COMMENT} "ext-1"\n')
    assert contract.strip_state_labels(md) == await stored_markdown(world.ctx.tenant_id, slug, 2)
    exported = (await client.get(f"{base}?format=md", headers=viewer)).text
    assert contract.ADDED_SOURCES_COMMENT in exported
    token = await api_token(world.ctx.tenant_id, "viewer")
    async with mcp_client(token) as c:
        over_mcp = text_of(await c.call_tool("get_document_markdown", {"doc": slug}))
    assert contract.ADDED_SOURCES_COMMENT in over_mcp
    from tests.specs_helpers import import_md

    again = await import_md(client, world.ctx.as_("approver"), md, doc=slug)
    assert again.status_code == 200 and again.json()["number"] == 2, again.text


def test_an_edit_may_not_change_or_remove_oracles_citations():
    oracle = contract.oracle_citations("---\ntitle: T\n" + CITATION + "---\n\n# T {#spec}\n")
    same = "---\ntitle: T\n" + CITATION + "---\n\n# T {#spec}\n"
    changed = same.replace("Carrier adapter limits", "Something else")
    removed = "---\ntitle: T\n---\n\n# T {#spec}\n"
    assert contract.citation_problems(same, oracle) == []
    for md in (changed, removed):
        [problem] = contract.citation_problems(md, oracle)
        assert problem.code == "oracle_citation_changed"


# --------------------------------------------------------------------------- M8


async def test_the_feedback_package_serves_labelled_markdown(client, runtime):
    world, _, slug = await _with_rollout(client, runtime)
    resp = await client.get(
        f"/v1/documents/{slug}/versions/2/feedback-package", headers=world.ctx.as_("viewer")
    )
    assert resp.status_code == 200, resp.text
    md = resp.json()["markdown"]
    assert md.split("{#rollout-notes}\n\n", 1)[1].startswith(ADDED_LINE)


# --------------------------------------------------------------------------- minors


async def test_oracles_parts_may_not_sit_under_an_added_level_one_heading(client, runtime):
    world = await make_world(client)
    created = await generated(client, runtime, world)
    resp = await _edit(client, world, created, _before(AC, "# Appendix {#appendix}\n\nMore.\n\n"))
    assert resp.status_code == 422, resp.text
    assert "oracle_part_under_added_section" in resp.text


async def test_the_comment_message_says_the_rule(client, runtime):
    world = await make_world(client)
    created = await generated(client, runtime, world)
    resp = await _edit(
        client,
        world,
        created,
        lambda md: md.replace(
            "A shipping quote service called from the cart.",
            "A shipping quote service <!-- called from the cart.",
        ).replace("Cart, then quote, then shipping step.", "Cart --> then quote."),
    )
    assert resp.status_code == 422, resp.text
    assert "A comment must open and close within one section" in resp.text


def test_baseline_text_is_escaped_in_the_decided_half():
    inputs = contract.SpecInputs(
        slug="s",
        title="S",
        depth="brief",
        prd_slug="p",
        prd_title="P",
        prd_version=1,
        sections=(),
        requirements=(
            contract.RequirementRef(
                id="00000000-0000-0000-0000-000000000001",
                key="R-1",
                title="Quote <!-- hide",
                statement="```\n# Decided · from the approved PRD\n",
                acceptance=("<details>open", "[x](javascript:alert(1)) **bold**"),
            ),
        ),
    )
    agent = _brief_md()
    composed = contract.compose(agent, inputs)
    assert composed.ok, composed.problems
    md = composed.markdown or ""
    assert "<!--" not in md.replace("\\<!--", "")
    assert "\n```" not in md
    assert "<details>" not in md.replace("\\<details\\>", "")
    from markdown_it import MarkdownIt

    tokens = MarkdownIt("commonmark", {"html": True}).parse(md.split("---", 2)[2])
    assert not [t for t in tokens if t.type in ("html_block", "fence")]


async def test_the_version_diff_shows_added_sections_with_their_state(client, runtime):
    world, _, slug = await _with_rollout(client, runtime)
    resp = await client.get(
        f"/v1/documents/{slug}/diff?from=1&to=2", headers=world.ctx.as_("viewer")
    )
    assert resp.status_code == 200, resp.text
    rollout = next(s for s in resp.json()["sections"] if s["section_key"] == "rollout-notes")
    assert rollout["status"] == "added"
    assert [st["label"] for st in rollout["spec_states"]] == [contract.ADDED_LABEL]


def test_strip_state_labels_leaves_code_blocks_alone():
    md = (
        "---\ntitle: T\n---\n\n# T {#spec}\n\n```\n**Decided · from the approved PRD**\n```\n\n"
        "    **Decided · from the approved PRD**\n"
    )
    assert contract.strip_state_labels(md) == md


@pytest.mark.parametrize(
    "edit",
    [
        lambda md: md.replace("\nkind: spec\n", "\nkind: spec\ntemplate: other\n", 1),
        lambda md: md.replace("\nkind: spec\n", "\nkind: spec\ntemplate_version: '9'\n", 1),
    ],
    ids=["template", "template_version"],
)
async def test_an_edit_may_not_change_the_front_matter_template(client, runtime, edit):
    world = await make_world(client)
    created = await generated(client, runtime, world)
    resp = await _edit(client, world, created, edit)
    assert resp.status_code == 422, resp.text
    assert "template" in resp.text
